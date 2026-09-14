"""Excel 导出（返回 bytes，不直接写文件）"""
import io
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def export_financial_tables_to_bytes(fin_res):
    """导出 3 张财务表（成本、利润、现金流）"""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        fin_res['df_cost'].to_excel(writer,      index=False, sheet_name="总成本费用表")
        fin_res['df_profit'].to_excel(writer,    index=False, sheet_name="利润与利润分配表")
        fin_res['cash_flow_df'].to_excel(writer, index=False, sheet_name="现金流量表")
    buf.seek(0)
    return buf.getvalue()


def _build_renewable_sheet(wb, sheet_name, df_8760, col_name, is_wind=True):
    """在 wb 里插入一张资源特性分析表（含 3 张图表）"""
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name="Source Han Sans CN", size=10, bold=True, color="FFFFFF")
    font_body   = Font(name="Source Han Sans CN", size=9)
    fill_header = PatternFill(
        start_color="366092" if is_wind else "C0504D",
        end_color="366092"   if is_wind else "C0504D",
        fill_type="solid",
    )
    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin",  color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"),
    )

    ws["A1"] = "月份";  ws["B1"] = "月发电量占比"

    df_temp = df_8760.copy()
    df_temp['month'] = pd.date_range("2026-01-01", periods=8760, freq="h").month
    m_sums = df_temp.groupby('month')[col_name].sum()
    total_sum = df_temp[col_name].sum()
    m_ratios = (m_sums / (total_sum + 1e-9)).tolist()

    for m in range(1, 13):
        ws.cell(row=m + 1, column=1, value=f"{m}月")
        cell = ws.cell(row=m + 1, column=2, value=m_ratios[m - 1])
        cell.number_format = "0.000"

    ws["D1"] = "时刻"; ws["E1"] = "最小出力"; ws["F1"] = "平均出力"; ws["G1"] = "最大出力"

    df_temp['hour_of_day'] = (np.arange(8760) % 24) + 1
    h_stats = df_temp.groupby('hour_of_day')[col_name].agg(['min', 'mean', 'max'])
    for h in range(1, 25):
        ws.cell(row=h + 1, column=4, value=h)
        ws.cell(row=h + 1, column=5, value=round(h_stats.loc[h, 'min'],  3)).number_format = "0.00"
        ws.cell(row=h + 1, column=6, value=round(h_stats.loc[h, 'mean'], 3)).number_format = "0.00"
        ws.cell(row=h + 1, column=7, value=round(h_stats.loc[h, 'max'],  3)).number_format = "0.00"

    ws.cell(row=15, column=1, value="出力系数")
    ws.cell(row=15, column=2, value="累积电量")
    ws.cell(row=15, column=3, value="保证率")

    bins = [i * 0.05 for i in range(21)]
    vals = df_temp[col_name].values
    for idx, b in enumerate(bins):
        r = 16 + idx
        ws.cell(row=r, column=1, value=f"0 ~ {b:.2f}")
        guar_r = float(np.mean(vals >= b))
        cum_e  = float(np.sum(vals[vals <= b]) / (total_sum + 1e-9))
        cell_e = ws.cell(row=r, column=2, value=round(cum_e, 4))
        cell_g = ws.cell(row=r, column=3, value=round(guar_r, 4))
        cell_e.number_format = "0.00%"
        cell_g.number_format = "0.00%"

    for col_range in [(1, 2, 1, 13), (4, 7, 1, 25), (1, 3, 15, 36)]:
        min_c, max_c, min_r, max_r = col_range
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = font_body
                cell.alignment = align_center
                cell.border = thin_border
                if r == min_r:
                    cell.font = font_header
                    cell.fill = fill_header

    # 图表 1：月发电量占比
    chart1 = BarChart()
    chart1.type = "col"; chart1.style = 10
    chart1.title = f"{'风电' if is_wind else '光伏'}月发电量占比"
    chart1.y_axis.number_format = "0.000"
    chart1.legend = None
    chart1.add_data(Reference(ws, min_col=2, min_row=1, max_row=13), titles_from_data=True)
    chart1.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    chart1.dataLabels = DataLabelList(); chart1.dataLabels.showVal = True
    chart1.width, chart1.height = 13, 8.5
    ws.add_chart(chart1, "I2")

    # 图表 2：出力特性
    chart2 = LineChart()
    chart2.title = f"{'风电' if is_wind else '光伏'}出力特性"
    chart2.style = 13; chart2.y_axis.number_format = "0.00"
    chart2.add_data(Reference(ws, min_col=5, max_col=7, min_row=1, max_row=25), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=4, min_row=2, max_row=25))
    chart2.width, chart2.height = 14, 8.5
    ws.add_chart(chart2, "Q2")

    # 图表 3：保证率-累积电量
    chart3 = LineChart()
    chart3.title = f"{'风电' if is_wind else '光伏'}电力 - 保证率 - 累积电量曲线"
    chart3.style = 12; chart3.y_axis.number_format = "0.0%"
    chart3.add_data(Reference(ws, min_col=2, max_col=3, min_row=15, max_row=36), titles_from_data=True)
    chart3.set_categories(Reference(ws, min_col=1, min_row=16, max_row=36))
    chart3.width, chart3.height = 18, 11.5
    ws.add_chart(chart3, "E16")


def export_full_report_to_bytes(df_8760, cap_res, fin_eval):
    """完整 Excel 报告（资源特性 + 优化结果）"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _build_renewable_sheet(wb, "风电特性分析", df_8760, 'wind_norm', is_wind=True)
    _build_renewable_sheet(wb, "光伏特性分析", df_8760, 'pv_norm',   is_wind=False)

    ws_res = wb.create_sheet(title="优化及财务评估结果")
    ws_res.views.sheetView[0].showGridLines = True

    font_title = Font(name="Source Han Sans CN", size=11, bold=True, color="FFFFFF")
    font_body  = Font(name="Source Han Sans CN", size=10)
    fill_head  = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_c    = Alignment(horizontal="center", vertical="center")
    thin_b = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin",  color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"),
    )

    rows = [
        ["指标类别", "指标名称", "数值", "单位"],
        ["容量配置结果", "推荐光伏容量", round(cap_res['PV_kW'] / 10000.0, 2), "万kW"],
        ["容量配置结果", "推荐风电容量", round(cap_res['Wind_kW'] / 10000.0, 2), "万kW"],
        ["容量配置结果", "推荐储能功率", round(cap_res['BESS_kW'] / 10000.0, 2), "万kW"],
        ["容量配置结果", "推荐储能容量", round(cap_res['BESS_kWh'] / 10000.0, 2), "万kWh"],
        ["投资评估", "光伏静态投资", round(fin_eval['inv_pv_wan'], 2), "万元"],
        ["投资评估", "风电静态投资", round(fin_eval['inv_wind_wan'], 2), "万元"],
        ["投资评估", "储能静态投资", round(fin_eval['inv_bess_total_wan'], 2), "万元"],
        ["投资评估", "项目静态总投资", round(fin_eval['total_static_inv_wan'] / 10000.0, 4), "亿元"],
        ["投资评估", "建设期利息", round(fin_eval['interest_construction_wan'], 2), "万元"],
        ["投资评估", "铺底/流动资金", round(fin_eval['working_capital_wan'], 2), "万元"],
        ["投资评估", "项目总投资(含利息与流动资金)", round(fin_eval['total_project_inv_wan'] / 10000.0, 4), "亿元"],
        ["投资评估", "项目资本金", round(fin_eval['equity_wan'], 2), "万元"],
        ["财务收益指标", "全投资IRR(税前)", round(fin_eval['irr_pre_tax'] * 100.0, 2), "%"],
        ["财务收益指标", "全投资IRR(税后)", round(fin_eval['irr_post_tax'] * 100.0, 2), "%"],
        ["财务收益指标", "资本金IRR", round(fin_eval['irr_equity'] * 100.0, 2), "%"],
    ]
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_res.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body
            cell.alignment = align_c
            cell.border = thin_b
            if r_idx == 1:
                cell.font = font_title
                cell.fill = fill_head

    ws_res.column_dimensions['A'].width = 18
    ws_res.column_dimensions['B'].width = 28
    ws_res.column_dimensions['C'].width = 15
    ws_res.column_dimensions['D'].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()